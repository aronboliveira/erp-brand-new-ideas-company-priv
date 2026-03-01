#!/usr/bin/env python3
"""
Payroll Exporter Module.

Handles Excel export for payroll data with employee salary
information and payment status.
"""
import sys
from typing import Any, Dict, List

import pandas as pd
from openpyxl.styles import Border, Font, PatternFill, Side

from base_exporter import BaseExporter, parse_number, safe_get


class PayrollExporter(BaseExporter):
    """
    Exporter for Payroll data.

    Generates xlsx files with employee ID, status, name,
    salary, net salary, and month information.
    """

    HEADINGS: List[str] = [
        "Employee Id", "Status", "Employee Name",
        "Salary", "Net Salary", "Month"
    ]
    COLUMN_WIDTHS: Dict[str, int] = {
        "A": 15, "B": 12, "C": 25, "D": 15, "E": 15, "F": 12
    }
    HEADER_FILL_COLOR: str = "CCE5FF"

    def __init__(self) -> None:
        """Initialize the Payroll exporter."""
        super().__init__("PayrollExporter")

    def get_headings(self) -> List[str]:
        """
        Return column headings for payroll export.

        Returns:
            List of column header strings.
        """
        return self.HEADINGS

    def process_data(self) -> pd.DataFrame:
        """
        Process input JSON data into a DataFrame.

        Expects data with 'rows' key containing payroll records.

        Returns:
            DataFrame with processed payroll data.
        """
        self.logger.info("Processing payroll data")
        rows: List[Dict[str, Any]] = self.data.get("rows", [])
        if not rows:
            self.logger.warning("No rows found in input data")
            return pd.DataFrame(columns=self.HEADINGS)
        processed: List[Dict[str, Any]] = []
        for idx, row in enumerate(rows):
            try:
                status_val = safe_get(row, "status", 0)
                status_str = "Paid" if status_val == 1 else "UnPaid"
                processed.append({
                    "Employee Id": safe_get(row, "employee_id", ""),
                    "Status": status_str,
                    "Employee Name": safe_get(row, "employee_name", ""),
                    "Salary": parse_number(
                        safe_get(row, "gross_salary", 0),
                        default=0.0
                    ),
                    "Net Salary": parse_number(
                        safe_get(row, "net_payable", 0),
                        default=0.0
                    ),
                    "Month": safe_get(row, "salary_month", ""),
                })
            except (KeyError, TypeError) as e:
                self.logger.error(
                    "Row %d processing failed: %s - row_data=%s",
                    idx, str(e), str(row)[:100]
                )
                continue
        self.logger.info("Processed %d rows successfully", len(processed))
        return pd.DataFrame(processed)

    def _apply_header_style(self) -> None:
        """Apply custom header styling for payroll."""
        if self.sheet is None:
            return
        header_fill = PatternFill(
            start_color=self.HEADER_FILL_COLOR,
            end_color=self.HEADER_FILL_COLOR,
            fill_type="solid"
        )
        header_border = Border(
            left=Side(style="medium", color="333333"),
            right=Side(style="medium", color="333333"),
            top=Side(style="medium", color="333333"),
            bottom=Side(style="medium", color="333333"),
        )
        for col in range(1, len(self.HEADINGS) + 1):
            cell = self.sheet.cell(row=1, column=col)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.border = header_border

    def _apply_body_style(self, row_count: int) -> None:
        """Apply borders to body rows."""
        if self.sheet is None or row_count < 1:
            return
        h_border = Border(
            top=Side(style="hair", color="E0E0E0"),
            bottom=Side(style="hair", color="E0E0E0"),
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
        )
        for row in range(2, row_count + 2):
            for col in range(1, len(self.HEADINGS) + 1):
                self.sheet.cell(row=row, column=col).border = h_border

    def export(self) -> None:
        """
        Execute the full export workflow.

        Reads JSON from stdin, processes data, creates styled workbook,
        and outputs binary xlsx to stdout.
        """
        self.logger.info("Starting payroll export")
        try:
            self.read_stdin()
            df = self.process_data()
            self.create_workbook()
            if self.sheet is None:
                self.logger.error("Failed to create worksheet")
                return
            self.sheet.title = "Payroll"
            for col, val in enumerate(self.get_headings(), start=1):
                self.sheet.cell(row=1, column=col, value=val)
            self._apply_header_style()
            rows_written = self.dataframe_to_sheet(df, start_row=2, include_header=False)
            self._apply_body_style(rows_written)
            if rows_written > 0:
                symbol = self.data.get("currency_symbol", "")
                self.apply_currency_format(
                    col=4,
                    row_start=2,
                    row_end=rows_written + 1,
                    symbol=symbol,
                )
                self.apply_currency_format(
                    col=5,
                    row_start=2,
                    row_end=rows_written + 1,
                    symbol=symbol,
                )
            self.set_column_widths(self.COLUMN_WIDTHS)
            self.freeze_pane("A2")
            output_path = self.data.get("output_path")
            if output_path:
                self.output_to_file(output_path)
                print(output_path, file=sys.stderr)
            else:
                self.output_to_stdout()
            self.logger.info("Payroll export completed successfully")
        except (ValueError, KeyError) as e:
            self.logger.error("Export failed due to data error: %s", str(e))
            sys.exit(1)
        except IOError as e:
            self.logger.error("Export failed due to IO error: %s", str(e))
            sys.exit(2)


def main() -> None:
    """Entry point for payroll export."""
    exporter = PayrollExporter()
    exporter.export()


if __name__ == "__main__":
    main()
