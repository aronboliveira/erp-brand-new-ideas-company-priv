"""Tests for base_exporter.py — ExportStyle, BaseExporter, and utility functions."""

import json
import os
import tempfile
from datetime import datetime
from io import StringIO
from unittest.mock import patch

import pandas as pd
import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill

from base_exporter import (
    BaseExporter,
    ExportStyle,
    format_currency,
    format_date,
    safe_get,
)


# ── ExportStyle constants ──────────────────────────────────────────


class TestExportStyle:
    def test_header_font_is_bold_white(self):
        assert ExportStyle.HEADER_FONT.bold is True
        assert ExportStyle.HEADER_FONT.color.rgb == "00FFFFFF"

    def test_header_fill_is_solid(self):
        assert ExportStyle.HEADER_FILL.fill_type == "solid"
        assert ExportStyle.HEADER_FILL.start_color.rgb == "00001122"

    def test_body_fill_light(self):
        assert ExportStyle.BODY_FILL_LIGHT.fill_type == "solid"

    def test_thin_border_has_sides(self):
        assert ExportStyle.THIN_BORDER.left.style == "thin"
        assert ExportStyle.THIN_BORDER.right.style == "thin"
        assert ExportStyle.THIN_BORDER.top.style == "hair"

    def test_header_border_all_thin(self):
        for side in ("left", "right", "top", "bottom"):
            assert getattr(ExportStyle.HEADER_BORDER, side).style == "thin"

    def test_center_align(self):
        assert ExportStyle.CENTER_ALIGN.horizontal == "center"
        assert ExportStyle.CENTER_ALIGN.vertical == "center"

    def test_left_align(self):
        assert ExportStyle.LEFT_ALIGN.horizontal == "left"


# ── Utility functions ──────────────────────────────────────────────


class TestSafeGet:
    def test_dict_key_exists(self):
        assert safe_get({"a": 1}, "a") == 1

    def test_dict_key_missing_returns_default(self):
        assert safe_get({"a": 1}, "b", "X") == "X"

    def test_none_returns_default(self):
        assert safe_get(None, "k", 42) == 42

    def test_object_attribute(self):
        class Obj:
            x = 10
        assert safe_get(Obj(), "x") == 10

    def test_object_missing_attr(self):
        assert safe_get(object(), "no_attr", "N/A") == "N/A"


class TestFormatDate:
    def test_none_returns_empty(self):
        assert format_date(None) == ""

    def test_iso_string(self):
        assert format_date("2025-06-15T10:30:00") == "2025-06-15"

    def test_iso_with_z(self):
        assert format_date("2025-06-15T00:00:00Z") == "2025-06-15"

    def test_datetime_object(self):
        dt = datetime(2025, 1, 1, 12, 0)
        assert format_date(dt) == "2025-01-01"

    def test_custom_format(self):
        assert format_date("2025-06-15T10:00:00", fmt="%d/%m/%Y") == "15/06/2025"

    def test_non_parsable_returns_as_is(self):
        assert format_date("not-a-date") == "not-a-date"

    def test_integer_returns_str(self):
        assert format_date(12345) == "12345"


class TestFormatCurrency:
    def test_none_returns_zero(self):
        assert format_currency(None) == "0.00"

    def test_integer(self):
        assert format_currency(1000) == "1,000.00"

    def test_float(self):
        assert format_currency(1234.5) == "1,234.50"

    def test_with_symbol(self):
        assert format_currency(99.9, symbol="$") == "$99.90"

    def test_custom_decimals(self):
        assert format_currency(5, symbol="€", decimals=0) == "€5"

    def test_non_numeric_returns_str(self):
        assert format_currency("abc") == "abc"


# ── BaseExporter (via concrete stub) ──────────────────────────────


class StubExporter(BaseExporter):
    """Minimal concrete exporter for testing abstract base methods."""

    HEADINGS = ["Col A", "Col B"]

    def get_headings(self):
        return self.HEADINGS

    def process_data(self):
        rows = self.data.get("rows", [])
        return pd.DataFrame(rows) if rows else pd.DataFrame(columns=self.HEADINGS)

    def export(self):
        self.read_stdin()
        df = self.process_data()
        self.create_workbook()
        for col, val in enumerate(self.get_headings(), 1):
            self.sheet.cell(row=1, column=col, value=val)
        self.dataframe_to_sheet(df, start_row=2, include_header=False)
        self.output_to_stdout()


class TestBaseExporterInit:
    def test_logger_name(self):
        exp = StubExporter("TestExp")
        assert exp.logger.name == "TestExp"

    def test_initial_state(self):
        exp = StubExporter("X")
        assert exp.data == {}
        assert exp.workbook is None
        assert exp.sheet is None


class TestReadStdin:
    def test_valid_json(self):
        payload = json.dumps({"rows": [{"a": 1}]})
        exp = StubExporter("T")
        with patch("sys.stdin", StringIO(payload)):
            result = exp.read_stdin()
        assert result == {"rows": [{"a": 1}]}
        assert exp.data == result

    def test_empty_stdin_raises(self):
        exp = StubExporter("T")
        with patch("sys.stdin", StringIO("")):
            with pytest.raises(ValueError, match="No input data"):
                exp.read_stdin()

    def test_invalid_json_raises(self):
        exp = StubExporter("T")
        with patch("sys.stdin", StringIO("{bad json")):
            with pytest.raises(json.JSONDecodeError):
                exp.read_stdin()


class TestCreateWorkbook:
    def test_creates_workbook_and_sheet(self):
        exp = StubExporter("T")
        wb = exp.create_workbook()
        assert isinstance(wb, Workbook)
        assert exp.sheet is not None
        assert exp.workbook is wb


class TestApplyHeaderStyle:
    def test_styles_applied_to_header_row(self):
        exp = StubExporter("T")
        exp.create_workbook()
        exp.sheet.cell(row=1, column=1, value="H1")
        exp.sheet.cell(row=1, column=2, value="H2")
        exp.apply_header_style(row_num=1, col_start=1, col_end=2)
        c1 = exp.sheet.cell(row=1, column=1)
        assert c1.font.bold is True
        assert c1.fill.start_color.rgb == "00001122"

    def test_no_sheet_does_not_crash(self):
        exp = StubExporter("T")
        exp.apply_header_style()  # sheet is None — should log warning but not crash


class TestApplyBodyStyle:
    def test_alternate_fill(self):
        exp = StubExporter("T")
        exp.create_workbook()
        for r in range(2, 6):
            exp.sheet.cell(row=r, column=1, value=f"r{r}")
        exp.apply_body_style(row_start=2, row_end=5, col_start=1, col_end=1)
        # Even rows get BODY_FILL_LIGHT
        assert exp.sheet.cell(row=2, column=1).fill.start_color.rgb == "00F5F5F5"
        # Odd rows should NOT have that fill
        assert exp.sheet.cell(row=3, column=1).fill.start_color.rgb != "00F5F5F5"


class TestFreezePane:
    def test_default_freeze(self):
        exp = StubExporter("T")
        exp.create_workbook()
        exp.freeze_pane()
        assert exp.sheet.freeze_panes == "A2"

    def test_custom_freeze(self):
        exp = StubExporter("T")
        exp.create_workbook()
        exp.freeze_pane("B3")
        assert exp.sheet.freeze_panes == "B3"


class TestSetColumnWidths:
    def test_widths_applied(self):
        exp = StubExporter("T")
        exp.create_workbook()
        exp.set_column_widths({"A": 20, "B": 30})
        assert exp.sheet.column_dimensions["A"].width == 20
        assert exp.sheet.column_dimensions["B"].width == 30


class TestDataframeToSheet:
    def test_writes_rows(self):
        exp = StubExporter("T")
        exp.create_workbook()
        df = pd.DataFrame({"Col A": [1, 2], "Col B": [3, 4]})
        written = exp.dataframe_to_sheet(df, start_row=1)
        assert written == 3  # header row + 2 data rows
        assert exp.sheet.cell(row=1, column=1).value == "Col A"
        assert exp.sheet.cell(row=2, column=1).value == 1


class TestOutputToFile:
    def test_saves_to_disk(self):
        exp = StubExporter("T")
        exp.create_workbook()
        exp.sheet.cell(row=1, column=1, value="test")
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            result = exp.output_to_file(path)
            assert os.path.exists(result)
            assert os.path.getsize(result) > 0
        finally:
            os.unlink(path)


class TestOutputToStdout:
    def test_writes_binary(self):
        exp = StubExporter("T")
        exp.create_workbook()
        exp.sheet.cell(row=1, column=1, value="test")
        import io

        buf = io.BytesIO()

        class FakeStdout:
            buffer = buf

        with patch("sys.stdout", FakeStdout()):
            exp.output_to_stdout()
        assert buf.tell() > 0


# ── Performance sanity ─────────────────────────────────────────────


class TestBaseExporterPerformance:
    @pytest.mark.timeout(5)
    def test_dataframe_to_sheet_100_rows(self):
        exp = StubExporter("Perf")
        exp.create_workbook()
        df = pd.DataFrame({"Col A": range(100), "Col B": range(100, 200)})
        written = exp.dataframe_to_sheet(df, start_row=1)
        assert written == 101  # header + 100

    @pytest.mark.timeout(5)
    def test_styling_100_rows(self):
        exp = StubExporter("Perf")
        exp.create_workbook()
        for r in range(1, 102):
            exp.sheet.cell(row=r, column=1, value=r)
        exp.apply_header_style(row_num=1, col_start=1, col_end=1)
        exp.apply_body_style(row_start=2, row_end=101, col_start=1, col_end=1)
