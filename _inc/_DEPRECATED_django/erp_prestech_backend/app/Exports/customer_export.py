import logging
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status

logger = logging.getLogger(__name__)
class CustomerExport:
  def __init__(self, customers_data, company_name='ERP Inc'):
    self.customers_data = customers_data
    self.company_name = company_name
    self.formatted_data = []
    self.df = pd.DataFrame()

  def format_data(self):
    try:
      temp = []
      for cust in self.customers_data:
        try:
          temp.append({
            "Customer No": cust.get("customer_id",""),
            "Name": cust.get("name",""),
            "Email": cust.get("email",""),
            "Contact": cust.get("contact",""),
            "Billing Name": cust.get("billing_name",""),
            "Billing Country": cust.get("billing_country",""),
            "Billing State": cust.get("billing_state",""),
            "Billing City": cust.get("billing_city",""),
            "Billing Phone": cust.get("billing_phone",""),
            "Billing Zip": cust.get("billing_zip",""),
            "Billing Address": cust.get("billing_address",""),
            "Shipping Name": cust.get("shipping_name",""),
            "Shipping Country": cust.get("shipping_country",""),
            "Shipping State": cust.get("shipping_state",""),
            "Shipping City": cust.get("shipping_city",""),
            "Shipping Phone": cust.get("shipping_phone",""),
            "Shipping Zip": cust.get("shipping_zip",""),
            "Shipping Address": cust.get("shipping_address",""),
            "Balance": cust.get("balance","")
          })
        except KeyError as e:
          logger.error(f"KeyError in format_data: {str(e)}")
        except Exception as e:
          logger.error(f"Unexpected error formatting customer: {str(e)}")
      self.formatted_data = temp
    except Exception as e:
      logger.error(f"format_data() error: {str(e)}")

  def dataframe(self):
    try:
      if not self.formatted_data:
        logger.warning("No data to convert into DataFrame.")
        return pd.DataFrame()
      self.df = pd.DataFrame(self.formatted_data)
      self.df.replace({np.nan:""}, inplace=True)
    except Exception as e:
      logger.error(f"dataframe() error: {str(e)}")
    return self.df

  def export_to_excel(self, file_name="customer_export.xlsx"):
    try:
      if self.df.empty:
        logger.warning("DataFrame is empty skipping Excel creation.")
        return None
      wb = Workbook()
      ws = wb.active
      ws.title = "Customers"

      for r in dataframe_to_rows(self.df, index=False, header=True):
        ws.append(r)

      bold_font = Font(bold=True)
      cold_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
      row_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
      thin_side_h = Side(border_style="thin", color="D9D9D9")
      thin_side_v = Side(border_style="thin", color="808080")

      for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
          cell.font = bold_font
          cell.fill = cold_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h, left=thin_side_v, right=thin_side_v)

      ws.freeze_panes = ws["A2"]

      for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
          cell.fill = row_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h, left=thin_side_v, right=thin_side_v)

      cols_count = len(self.df.columns)
      if cols_count>0:
        top_cell = f"A1:{chr(64+cols_count)}1"
        ws.merge_cells(top_cell)
        ws["A1"] = f"Customer Export - {self.company_name}"
        ws["A1"].font = bold_font
        ws["A1"].alignment = Alignment(horizontal="center")

      wb.save(file_name)
      return file_name
    except PermissionError as e:
      logger.error(f"Permission error: {str(e)}")
      return None
    except Exception as e:
      logger.error(f"export_to_excel() error: {str(e)}")
      return None

class CustomerExportAPIView(APIView):
  def get(self, request, *args, **kwargs):
    try:
      dummy_customers = [
        {
          "customer_id": "CUST-001",
          "name": "John Doe",
          "email": "john@example.com",
          "contact": "+123456789",
          "billing_name": "John Inc",
          "billing_country": "USA",
          "billing_state": "CA",
          "billing_city": "Los Angeles",
          "billing_phone": "+123456789",
          "billing_zip": "90001",
          "billing_address": "Sunset Blvd 123",
          "shipping_name": "John Inc S",
          "shipping_country": "USA",
          "shipping_state": "CA",
          "shipping_city": "Los Angeles",
          "shipping_phone": "+123456789",
          "shipping_zip": "90002",
          "shipping_address": "Sunset Blvd 124",
          "balance": "50.00"
        },
        {
          "customer_id": "CUST-002",
          "name": "Jane Smith",
          "email": "jane@example.com",
          "contact": "+987654321",
          "billing_name": "Jane LLC",
          "billing_country": "USA",
          "billing_state": "NY",
          "billing_city": "New York",
          "billing_phone": "+987654321",
          "billing_zip": "10001",
          "billing_address": "Times Square 10",
          "shipping_name": "Jane LLC S",
          "shipping_country": "USA",
          "shipping_state": "NY",
          "shipping_city": "New York",
          "shipping_phone": "+987654321",
          "shipping_zip": "10002",
          "shipping_address": "Times Square 11",
          "balance": "120.00"
        }
      ]
      exporter = CustomerExport(dummy_customers, company_name="ERP Inc")
      exporter.format_data()
      exporter.dataframe()
      filename = exporter.export_to_excel()
      if not filename:
        return Response({"error": "Failed to generate Excel"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
      with open(filename,"rb") as f:
        content = f.read()
      resp = HttpResponse(content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
      resp["Content-Disposition"] = f'attachment filename="{filename}"'
      return resp
    except FileNotFoundError as e:
      logger.error(f"File not found: {str(e)}")
      return Response({"error":"File not found."}, status=status.HTTP_404_NOT_FOUND)
    except PermissionError as e:
      logger.error(f"Permission Error: {str(e)}")
      return Response({"error":"Permission Error."}, status=status.HTTP_403_FORBIDDEN)
    except Exception as e:
      logger.error(f"Unexpected error: {str(e)}")
      return Response({"error":"Unexpected error occurred."}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
