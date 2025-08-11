import logging
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

logger = logging.getLogger(__name__)

class ProposalExport:
    def __init__(self, proposals_data, company_name='ERP Inc'):
        self.proposals_data = proposals_data
        self.company_name = company_name
        self.formatted_data = []
        self.df = pd.DataFrame()
        
    def format_data(self):
        try:
            temp = []
            for proposal in self.proposals_data:
                try:
                    proposal_id = proposal.get('proposal_id', '')
                    issue_date = proposal.get('issue_date', '')
                    send_date = proposal.get('send_date', '')
                    category_id = proposal.get('category_id', 'Income')  # Placeholder for ProductServiceCategory lookup
                    status_index = proposal.get('status', 0)  # Assuming int index, adjust accordingly
                    status = self.statuses().get(status_index, 'Unknown')

                    temp.append({
                        "ID": proposal_id,
                        "Proposal No": proposal_id,  # Simulated proposalNumberFormat
                        "Issue Date": issue_date,
                        "Send Date": send_date,
                        "Category": category_id,
                        "Status": status
                    })

                except KeyError as e:
                    logger.error(f"KeyError in format_data for proposal: {str(e)}")
                except Exception as e:
                    logger.error(f"Unexpected error formatting proposal: {str(e)}")

            self.formatted_data = temp
        except Exception as e:
            logger.error(f"format_data() error: {str(e)}")

    def statuses(self):
        # Simulate Proposal::$statues
        return {
            0: 'Draft',
            1: 'Sent',
            2: 'Accepted',
            3: 'Declined',
            4: 'Expired'
        }

    def dataframe(self):
        try:
            if not self.formatted_data:
                logger.warning("No data to build DataFrame.")
                return pd.DataFrame()

            self.df = pd.DataFrame(self.formatted_data)
            self.df.replace({np.nan: ""}, inplace=True)
        except Exception as e:
            logger.error(f"dataframe() error: {str(e)}")

        return self.df

    def export_to_excel(self, filename="proposal_export.xlsx"):
        try:
            if self.df.empty:
                logger.warning("DataFrame is empty no Excel to generate.")
                return None

            wb = Workbook()
            ws = wb.active
            ws.title = "Proposals"

            for r in dataframe_to_rows(self.df, index=False, header=True):
                ws.append(r)

            bold_font = Font(bold=True)
            cold_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
            row_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            thin_side_h = Side(border_style="thin", color="D9D9D9")
            thin_side_v = Side(border_style="thin", color="808080")

            for col in ws.iter_cols(min_row=1, max_row=1):
                for cell in col:
                    cell.font = bold_font
                    cell.fill = cold_fill
                    cell.border = Border(top=thin_side_h, bottom=thin_side_h,
                                         left=thin_side_v, right=thin_side_v)

            ws.freeze_panes = ws["A2"]

            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for cell in row:
                    cell.fill = row_fill
                    cell.border = Border(top=thin_side_h, bottom=thin_side_h,
                                         left=thin_side_v, right=thin_side_v)

            cols_count = len(self.df.columns)
            if cols_count > 0:
                top_cell = f"A1:{chr(64 + cols_count)}1"
                ws.merge_cells(top_cell)
                ws["A1"] = f"Proposal Export - {self.company_name}"
                ws["A1"].font = bold_font
                ws["A1"].alignment = Alignment(horizontal="center")

            wb.save(filename)
            return filename

        except PermissionError as e:
            logger.error(f"Permission error: {str(e)}")
            return None
        except Exception as e:
            logger.error(f"export_to_excel() error: {str(e)}")
            return None
