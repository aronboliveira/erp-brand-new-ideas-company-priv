import logging
import pandas as pd
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime
from colorama import Fore
logger = logging.getLogger(__name__)
class BalanceSheetExport:
  def __init__(self, data, start_date, end_date, company_name):
    self.raw_data = data
    self.start_date = start_date
    self.end_date = end_date
    self.company_name = company_name
    self.formatted_data = []
    self.df = None
  def format_data(self):
    try:
      fd = []
      loe_enc = False
      loe_enc1 = False
      liab = False
      assets = False
      total = 0
      for cat, subcats in self.raw_data.items():
        amt_total = 0
        for sub in subcats:
          for a in sub.get('account', []):
            try:
              net = a.get('netAmount')
              if net is not None:
                if cat in ['Liabilities','Equity']:
                  if not loe_enc:
                    fd.append({'Account Name': 'Liabilities & Equity','Account No': '','Total': ''})
                    loe_enc = True
                  if not liab:
                    fd.append({'Account Name': f'  {cat}','Account No': '','Total': ''})
                    liab = True
                else:
                  if not assets:
                    fd.append({'Account Name': cat,'Account No': '','Total': ''})
                    assets = True
              break
            except Exception as e:
              logger.error(f"Error processing account in category {cat}: {str(e)}")
        for sub in subcats:
          for a in sub.get('account', []):
            try:
              net = a.get('netAmount')
              if net is not None:
                fd.append({'Account Name': f'    {sub.get("subType")}','Account No': '','Total': ''})
                break
            except Exception as e:
              logger.error(f"Error processing sub_category {sub.get('subType')}: {str(e)}")
          for a in sub.get('account', []):
            try:
              net = a.get('netAmount')
              if net is not None:
                fd.append({'Account Name': f'       {a.get("account_name")}','Account No': a.get('account_no'),'Total': net})
                amt_total += net if isinstance(net, (int,float)) else 0
            except Exception as e:
              logger.error(f"Error processing account {a.get('account_name')}: {str(e)}")
          for a in sub.get('account', []):
            try:
              net = a.get('netAmount')
              if net is not None:
                fd.append({'Account Name': f'    Total {sub.get("subType")}','Account No': '','Total': amt_total})
                break
            except Exception as e:
              logger.error(f"Error processing total for sub_category {sub.get('subType')}: {str(e)}")
        if cat in ['Liabilities','Equity'] and amt_total != 0:
          fd.append({'Account Name': f'  Total {cat}','Account No': '','Total': amt_total})
        elif amt_total != 0:
          fd.append({'Account Name': f'Total {cat}','Account No': '','Total': amt_total})
      for a in fd:
        try:
          an = a.get('Account Name')
          total += a.get('Total',0) if an in ['  Total Liabilities','  Total Equity'] else 0
        except Exception as e:
          logger.error(f"Error summing totals: {str(e)}")
      if not loe_enc1:
        fd.append({'Account Name': 'Total Liabilities & Equity','Account No': '','Total': total})
        loe_enc1 = True
      self.formatted_data = fd
    except Exception as e:
      logger.error(f"format_data() error: {str(e)}")
  def dataframe(self):
    try:
      if not self.formatted_data:
        return pd.DataFrame()
      self.df = pd.DataFrame(self.formatted_data)
      self.df.replace({pd.NA: ''}, inplace=True)
    except Exception as e:
      logger.error(f"dataframe() error: {str(e)}")
    return self.df
  def export_to_excel(self, file_name="balance_sheet.xlsx"):
    try:
      if self.df is None or self.df.empty:
        return
      wb = Workbook()
      ws = wb.active
      ws.title = "Balance Sheet"
      for r in dataframe_to_rows(self.df, index=False, header=True):
        ws.append(r)
      bold_font = Font(bold=True)
      cold_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
      row_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
      thin_hor = Side(border_style="thin", color="D9D9D9")
      thin_ver = Side(border_style="thin", color="808080")
      for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
          cell.font = bold_font
          cell.fill = cold_fill
          cell.border = Border(top=thin_hor, bottom=thin_hor, left=thin_ver, right=thin_ver)
      ws.freeze_panes = ws["A2"]
      for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
          cell.fill = row_fill
          cell.border = Border(top=thin_hor, bottom=thin_hor, left=thin_ver, right=thin_ver)
      ws.merge_cells("A1:F1")
      ws.merge_cells("A2:F2")
      ws.merge_cells("A3:F3")
      ws["A1"] = f"Balance Sheet - {self.company_name}"
      ws["A2"] = f"Print Out Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
      ws["A3"] = f"Date: {self.start_date} - {self.end_date}"
      ws["A1"].font = bold_font 
      ws["A2"].font = bold_font 
      ws["A3"].font = bold_font
      ws["A1"].alignment = Alignment(horizontal="center")
      ws["A2"].alignment = Alignment(horizontal="center")
      ws["A3"].alignment = Alignment(horizontal="center")
      wb.save(file_name)
      print(Fore.GREEN + f"Excel exported: {file_name}")
    except Exception as e:
      logger.error(f"export_to_excel() error: {str(e)}")
  def export_to_pdf(self, file_name="balance_sheet.pdf", image_folder="images/"):
    try:
      from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image, PageBreak
      from reportlab.lib import colors
      from reportlab.lib.pagesizes import A4
      from reportlab.lib.styles import getSampleStyleSheet
      doc = SimpleDocTemplate(file_name, pagesize=A4)
      styles = getSampleStyleSheet()
      elems = []
      elems.append(Paragraph(f"Balance Sheet - {self.company_name}", styles['Title']))
      elems.append(Paragraph(f"Print Out Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
      elems.append(Paragraph(f"Date: {self.start_date} - {self.end_date}", styles['Normal']))
      elems.append(Spacer(1,12))
      if self.df is None or self.df.empty:
        elems.append(Paragraph("No Data Available", styles['Normal']))
      else:
        hdrs = list(self.df.columns)
        rows = [hdrs] + self.df.values.tolist()
        tbl = Table(rows)
        tbl.setStyle(TableStyle([
          ('BACKGROUND',(0,0),(-1,0),colors.HexColor("#003366")),
          ('TEXTCOLOR',(0,0),(-1,0),colors.white),
          ('ALIGN',(0,0),(-1,-1),'CENTER'),
          ('GRID',(0,0),(-1,-1),0.25,colors.grey),
          ('BACKGROUND',(0,1),(-1,-1),colors.whitesmoke)
        ]))
        elems.append(tbl)
      elems.append(PageBreak())
      elems.append(Paragraph("Attached Images", styles['Heading2']))
      if os.path.exists(image_folder):
        for img in os.listdir(image_folder):
          if img.lower().endswith((".png",".jpg",".jpeg")):
            img_path = os.path.join(image_folder, img)
            elems.append(Image(img_path, width=200, height=150))
            elems.append(Spacer(1,12))
      else:
        elems.append(Paragraph("No Images Found", styles['Normal']))
      doc.build(elems)
      print(Fore.GREEN + f"PDF exported: {file_name}")
    except Exception as e:
      logger.error(f"export_to_pdf() error: {str(e)}")
  def detect_tables(self):
    try:
      tables = {}
      if self.df is None or self.df.empty:
        return tables
      df_rows = self.df.iterrows()
      header_idx = None
      headers = ()
      for idx, row in df_rows:
        vals = row.values.tolist()
        cnt = sum(1 for v in vals if v)
        if cnt >= 2:
          header_idx = idx
          headers = tuple(v for v in vals if v)
          break
      if not headers:
        return {}
      col_start = self.df.columns.get_loc(headers[0])
      table_dict = {}
      for i in range(header_idx+1, len(self.df)):
        rec = {}
        for j, h in enumerate(headers):
          try:
            val = self.df.iloc[i, col_start+j] if (col_start+j) < len(self.df.columns) else ''
            rec[h] = '' if pd.isna(val) else val
          except Exception as e:
            logger.error(f"Cell read error at row {i} col {j}: {str(e)}")
        table_dict[i] = rec
      tables[0] = {'header_row': header_idx, 'headers': headers, 'rows': table_dict}
      return tables
    except Exception as e:
      logger.error(f"detect_tables() error: {str(e)}")
      return {}
  def ask_to_save_json(self, tables_dict):
    try:
      if not tables_dict:
        return
      inp = input("Save result as JSON? (y/n): ")
      inp = 'y' if inp.strip().lower()=='y' else 'n'
      if inp=='y':
        try:
          with open('balance_sheet.json','w',encoding='utf-8') as f:
            json.dump(tables_dict, f, ensure_ascii=False, indent=2)
          print(Fore.GREEN + "JSON saved: balance_sheet.json")
        except Exception as e:
          logger.error(f"Error saving JSON file: {str(e)}")
    except Exception as e:
      logger.error(f"ask_to_save_json() error: {str(e)}")
  def ask_to_export_pdf(self):
    try:
      inp = input("Export PDF? (y/n): ")
      inp = 'y' if inp.strip().lower()=='y' else 'n'
      if inp=='y':
        self.export_to_pdf()
    except Exception as e:
      logger.error(f"ask_to_export_pdf() error: {str(e)}")
  def run(self):
    try:
      print(Fore.CYAN + "Export process started...")
      self.format_data()
      self.dataframe()
      self.export_to_excel()
      self.ask_to_export_pdf()
      tbls = self.detect_tables()
      self.ask_to_save_json(tbls)
      print(Fore.GREEN + "Export process completed.")
    except Exception as e:
      logger.error(f"run() error: {str(e)}")

