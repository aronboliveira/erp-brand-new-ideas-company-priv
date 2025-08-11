import logging
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

logger = logging.getLogger(__name__)

class BillExport:
  def __init__(self, bills_data, company_name='ERP Inc'):
    self.bills_data = bills_data
    self.company_name = company_name
    self.formatted_data = []
    self.df = pd.DataFrame()

  def format_data(self):
    try:
      formatted_data = []
      for bill in self.bills_data:
        try:
          bill_no = bill.get('bill_id') or ''
          bill_date = bill.get('bill_date') or ''
          due_date = bill.get('due_date') or ''
          order_no = bill.get('order_no') or ''
          status = bill.get('status') or ''
          send_date = bill.get('send_date') or ''
          category = bill.get('category') or ''

          formatted_data.append({
            'Bill No': bill_no,
            'Bill Date': bill_date,
            'Due Date': due_date,
            'Order No': order_no,
            'Status': status,
            'Send Date': send_date,
            'Category': category
          })
        except KeyError as e:
          logger.error(f'Missing key in bill record: {str(e)}')
        except Exception as e:
          logger.error(f'Unexpected error formatting bill: {str(e)}')

      self.formatted_data = formatted_data
    except Exception as e:
      logger.error(f'format_data() error: {str(e)}')

  def dataframe(self):
    try:
      if not self.formatted_data:
        logger.warning("No data to convert to DataFrame.")
        return pd.DataFrame()
      self.df = pd.DataFrame(self.formatted_data)
      self.df.replace({np.nan: ''}, inplace=True)
      return self.df
    except Exception as e:
      logger.error(f'dataframe() error: {str(e)}')
      return pd.DataFrame()

  def export_to_excel(self, filename='bill_export.xlsx'):
    try:
      if self.df.empty:
        logger.warning("DataFrame is empty skipping Excel generation.")
        return None

      wb = Workbook()
      ws = wb.active
      ws.title = "Bills"

      for r in dataframe_to_rows(self.df, index=False, header=True):
        ws.append(r)

      bold_font = Font(bold=True)
      cold_fill = PatternFill(start_color='003366', end_color='003366', fill_type='solid')
      row_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
      thin_border_horizontal = Side(border_style="thin", color="D9D9D9")
      thin_border_vertical = Side(border_style="thin", color="808080")

      for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
          cell.font = bold_font
          cell.fill = cold_fill
          cell.border = Border(top=thin_border_horizontal, bottom=thin_border_horizontal,
                               left=thin_border_vertical, right=thin_border_vertical)

      ws.freeze_panes = ws['A2']

      for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
          cell.fill = row_fill
          cell.border = Border(top=thin_border_horizontal, bottom=thin_border_horizontal,
                               left=thin_border_vertical, right=thin_border_vertical)

      ws.merge_cells('A1:G1')
      ws['A1'] = f'Bill Export - {self.company_name}'
      ws['A1'].font = bold_font
      ws['A1'].alignment = Alignment(horizontal="center")

      wb.save(filename)
      return filename
    except PermissionError as e:
      logger.error(f'Permission denied: {str(e)}')
      return None
    except Exception as e:
      logger.error(f'export_to_excel() error: {str(e)}')
      return None
