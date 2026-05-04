from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from conftest import load_exporter_class


WORKBOOK_CASES = [
    ("account_statement_exporter", "Account Statement", "Account Statement", 1),
    ("balance_sheet_exporter", "Balance Sheet", "Dashboard", 5),
    ("customer_exporter", "Customers", "Customers", 1),
    ("employee_exporter", "Employees", "Employees", 1),
    ("payroll_exporter", "Payroll", "Payroll", 1),
    ("payslip_exporter", "Payslips", "Payslips", 1),
    ("product_service_exporter", "Products & Services", "Products & Services", 1),
    ("sales_report_exporter", "Sales Item", "Dashboard", 6),
    ("transaction_exporter", "Transactions", "Transactions", 1),
    ("trial_balance_exporter", "Trial Balance", "Dashboard", 6),
    ("vendor_exporter", "Vendors", "Vendors", 1),
]


def _run_export(module_name: str, payload: dict[str, Any], tmp_path: Path, monkeypatch: Any) -> Any:
    exporter_cls = load_exporter_class(module_name)
    output_path = tmp_path / f"{module_name}.xlsx"
    payload = dict(payload)
    payload["output_path"] = str(output_path)

    monkeypatch.setattr("sys.stdin", io.StringIO(json.dumps(payload)))
    monkeypatch.setattr("sys.stderr", io.StringIO())

    exporter = exporter_cls()
    exporter.export()
    return openpyxl.load_workbook(output_path, data_only=False)


def _sheet_values(sheet: Any) -> list[Any]:
    return [cell.value for row in sheet.iter_rows() for cell in row if cell.value not in (None, "")]


def _freeze_panes(sheet: Any) -> Any:
    pane = sheet.freeze_panes
    return pane.coordinate if hasattr(pane, "coordinate") else pane


@pytest.mark.parametrize(  # type: ignore[untyped-decorator]
    "module_name,main_sheet,first_sheet,header_row", WORKBOOK_CASES
)
def test_exporters_generate_loadable_workbooks(
    module_name: str,
    main_sheet: str,
    first_sheet: str,
    header_row: int,
    exporter_payloads: Any,
    tmp_path: Path,
    monkeypatch: Any,
) -> None:
    workbook = _run_export(module_name, exporter_payloads[module_name](), tmp_path, monkeypatch)

    assert workbook.sheetnames[0] == first_sheet
    assert main_sheet in workbook.sheetnames
    sheet = workbook[main_sheet]
    headers = [sheet.cell(row=header_row, column=index).value for index in range(1, 5)]

    if module_name == "sales_report_exporter":
        assert headers == ["Item Name", "Quantity Sold", "Amount", "Average Amount"]
        assert "Sales Pivot" in workbook.sheetnames
        assert _freeze_panes(sheet) == "A7"
    elif module_name == "balance_sheet_exporter":
        assert headers[:3] == ["Account", "Account No", "Total"]
        assert _freeze_panes(sheet) == "A6"
    elif module_name == "trial_balance_exporter":
        assert headers == ["Account Name", "Account No", "Debit", "Credit"]
        assert _freeze_panes(sheet) == "A7"
    else:
        assert headers[0] is not None
        assert _freeze_panes(sheet) == "A2"


def test_balance_sheet_workbook_contains_dashboard_and_balance_check(
    exporter_payloads: Any, tmp_path: Path, monkeypatch: Any
) -> None:
    workbook = _run_export(
        "balance_sheet_exporter",
        exporter_payloads["balance_sheet_exporter"](),
        tmp_path,
        monkeypatch,
    )

    dashboard = workbook["Dashboard"]
    sheet = workbook["Balance Sheet"]
    values = _sheet_values(sheet)

    assert dashboard["B2"].value == "Balance Sheet Dashboard - PrestTech"
    assert "Balance Check:" in values
    assert "Total Liabilities & Equity" in values


def test_trial_balance_workbook_contains_balance_formula(
    exporter_payloads: Any, tmp_path: Path, monkeypatch: Any
) -> None:
    workbook = _run_export(
        "trial_balance_exporter",
        exporter_payloads["trial_balance_exporter"](),
        tmp_path,
        monkeypatch,
    )

    dashboard = workbook["Dashboard"]
    sheet = workbook["Trial Balance"]
    formulas = [
        cell.value
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]

    assert dashboard["B4"].value == "✓ Trial Balance is BALANCED"
    assert any(formula.startswith('=IF(ABS(C') for formula in formulas)
    assert any(formula.startswith("=C") and "-D" in formula for formula in formulas)


def test_sales_report_workbook_contains_totals_and_pivot(
    exporter_payloads: Any, tmp_path: Path, monkeypatch: Any
) -> None:
    workbook = _run_export(
        "sales_report_exporter",
        exporter_payloads["sales_report_exporter"](),
        tmp_path,
        monkeypatch,
    )

    sheet = workbook["Sales Item"]
    pivot = workbook["Sales Pivot"]
    values = _sheet_values(sheet)
    formulas = [
        cell.value
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]

    assert "TOTAL" in values
    assert any(formula.startswith("=SUM(B7:B12)") for formula in formulas)
    assert any(formula.startswith("=SUM(C7:C12)") for formula in formulas)
    assert pivot["A1"].value is not None
