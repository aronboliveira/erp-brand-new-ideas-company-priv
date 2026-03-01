#!/usr/bin/env python3
"""
Vendor Exporter Module.

Handles Excel export for vendor data with contact
and address information.
"""
import sys
from typing import Any, Dict, List

import pandas as pd
from openpyxl.styles import Border, Font, PatternFill, Side

from base_exporter import BaseExporter, parse_number, safe_get


class VendorExporter(BaseExporter):
    """
    Exporter for Vendor data.

    Generates xlsx files with vendor ID, contact info,
    billing/shipping addresses, and balance.
    """

    HEADINGS: List[str] = [
        "ID", "Name", "Email", "Contact",
        "Billing Name", "Billing Country", "Billing State", "Billing City",
        "Billing Phone", "Billing Zip", "Billing Address",
        "Shipping Name", "Shipping Country", "Shipping State", "Shipping City",
        "Shipping Phone", "Shipping Zip", "Shipping Address", "Balance"
    ]
    COLUMN_WIDTHS: Dict[str, int] = {
        "A": 12, "B": 20, "C": 25, "D": 15,
        "E": 18, "F": 15, "G": 15, "H": 15,
        "I": 15, "J": 12, "K": 30,
        "L": 18, "M": 15, "N": 15, "O": 15,
        "P": 15, "Q": 12, "R": 30, "S": 15
    }
    HEADER_FILL_COLOR: str = "001122"

    def __init__(self) -> None:
        """Initialize the Vendor exporter."""
        super().__init__("VendorExporter")

    def get_headings(self) -> List[str]:
        """
        Return column headings for vendor export.

        Returns:
            List of column header strings.
        """
        return self.HEADINGS

    def process_data(self) -> pd.DataFrame:
        """
        Process input JSON data into a DataFrame.

        Expects data with 'rows' key containing vendor records.

        Returns:
            DataFrame with processed vendor data.
        """
        self.logger.info("Processing vendor data")
        rows: List[Dict[str, Any]] = self.data.get("rows", [])
        if not rows:
            self.logger.warning("No rows found in input data")
            return pd.DataFrame(columns=self.HEADINGS)
        processed: List[Dict[str, Any]] = []
        for idx, row in enumerate(rows):
            try:
                processed.append({
                    "ID": safe_get(row, "vendor_no", ""),
                    "Name": safe_get(row, "name", ""),
                    "Email": safe_get(row, "email", ""),
                    "Contact": safe_get(row, "contact", ""),
                    "Billing Name": safe_get(row, "billing_name", ""),
                    "Billing Country": safe_get(row, "billing_country", ""),
                    "Billing State": safe_get(row, "billing_state", ""),
                    "Billing City": safe_get(row, "billing_city", ""),
                    "Billing Phone": safe_get(row, "billing_phone", ""),
                    "Billing Zip": safe_get(row, "billing_zip", ""),
                    "Billing Address": safe_get(row, "billing_address", ""),
                    "Shipping Name": safe_get(row, "shipping_name", ""),
                    "Shipping Country": safe_get(row, "shipping_country", ""),
                    "Shipping State": safe_get(row, "shipping_state", ""),
                    "Shipping City": safe_get(row, "shipping_city", ""),
                    "Shipping Phone": safe_get(row, "shipping_phone", ""),
                    "Shipping Zip": safe_get(row, "shipping_zip", ""),
                    "Shipping Address": safe_get(row, "shipping_address", ""),
                    "Balance": parse_number(
                        safe_get(row, "balance", 0),
                        default=0.0
                    ),
                })
            except (KeyError, TypeError) as e:
                self.logger.error(
                    "Row %d processing failed: %s - row_data=%s",
                    idx, str(e), str(row)[:100]
                )
                continue
        self.logger.info("Processed %d rows successfully", len(processed))
        return pd.DataFrame(processed)

    def _apply_header_style(self) -> None:
        """Apply custom header styling for vendor export."""
        if self.sheet is None:
            return
        header_fill = PatternFill(
            start_color=self.HEADER_FILL_COLOR,
            end_color=self.HEADER_FILL_COLOR,
            fill_type="solid"
        )
        header_border = Border(
            left=Side(style="thin", color="333333"),
            right=Side(style="thin", color="333333"),
            top=Side(style="hair", color="333333"),
            bottom=Side(style="hair", color="333333"),
        )
        for col in range(1, len(self.HEADINGS) + 1):
            cell = self.sheet.cell(row=1, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.border = header_border

    def export(self) -> None:
        """
        Execute the full export workflow.

        Reads JSON from stdin, processes data, creates styled workbook,
        and outputs binary xlsx to stdout.
        """
        self.logger.info("Starting vendor export")
        try:
            self.read_stdin()
            df = self.process_data()
            self.create_workbook()
            if self.sheet is None:
                self.logger.error("Failed to create worksheet")
                return
            self.sheet.title = "Vendors"
            for col, val in enumerate(self.get_headings(), start=1):
                self.sheet.cell(row=1, column=col, value=val)
            self._apply_header_style()
            rows_written = self.dataframe_to_sheet(df, start_row=2, include_header=False)
            if rows_written > 0:
                self.apply_body_style(
                    row_start=2,
                    row_end=rows_written + 1,
                    col_start=1,
                    col_end=len(self.HEADINGS)
                )
                self.apply_currency_format(
                    col=19,
                    row_start=2,
                    row_end=rows_written + 1,
                    symbol=self.data.get("currency_symbol", ""),
                )
            self.set_column_widths(self.COLUMN_WIDTHS)
            self.freeze_pane("A2")
            output_path = self.data.get("output_path")
            if output_path:
                self.output_to_file(output_path)
                print(output_path, file=sys.stderr)
            else:
                self.output_to_stdout()
            self.logger.info("Vendor export completed successfully")
        except (ValueError, KeyError) as e:
            self.logger.error("Export failed due to data error: %s", str(e))
            sys.exit(1)
        except IOError as e:
            self.logger.error("Export failed due to IO error: %s", str(e))
            sys.exit(2)


def main() -> None:
    """Entry point for vendor export."""
    exporter = VendorExporter()
    exporter.export()


if __name__ == "__main__":
    main()
