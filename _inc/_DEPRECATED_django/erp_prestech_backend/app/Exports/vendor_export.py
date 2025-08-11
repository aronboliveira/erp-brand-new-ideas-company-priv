import logging
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

logger = logging.getLogger(__name__)

class VendorExport:
  def __init__(self, vendors_data, company_name='ERP Inc'):
    self.vendors_data = vendors_data
    self.company_name = company_name
    self.formatted_data = []
    self.df = pd.DataFrame()

  def format_data(self):
    try:
      temp = []
      # In a real scenario, you'd fetch from Vendor model with filters. We'll parse the passed array.
      for v in self.vendors_data:
        try:
          temp.append({
            "ID": v.get("vendor_id", ""),
            "Name": v.get("name", ""),
            "Email": v.get("email", ""),
            "Contact": v.get("contact", ""),
            "Billing Name": v.get("billing_name", ""),
            "Billing Country": v.get("billing_country", ""),
            "Billing State": v.get("billing_state", ""),
            "Billing City": v.get("billing_city", ""),
            "Billing Phone": v.get("billing_phone", ""),
            "Billing Zip": v.get("billing_zip", ""),
            "Billing Address": v.get("billing_address", ""),
            "Shipping Name": v.get("shipping_name", ""),
            "Shipping Country": v.get("shipping_country", ""),
            "Shipping State": v.get("shipping_state", ""),
            "Shipping City": v.get("shipping_city", ""),
            "Shipping Phone": v.get("shipping_phone", ""),
            "Shipping Zip": v.get("shipping_zip", ""),
            "Shipping Address": v.get("shipping_address", ""),
            "Balance": v.get("balance", "")
          })
        except Exception as e:
          logger.error(f"Error formatting vendor: {str(e)}")

      self.formatted_data = temp
    except Exception as e:
      logger.error(f"format_data() error: {str(e)}")

  def dataframe(self):
    try:
      if not self.formatted_data:
        logger.warning("No vendor data found to create DataFrame.")
        return pd.DataFrame()
      self.df = pd.DataFrame(self.formatted_data)
      self.df.replace({np.nan: ""}, inplace=True)
    except Exception as e:
      logger.error(f"dataframe() error: {str(e)}")
    return self.df

  def export_to_excel(self, filename="vendor_export.xlsx"):
    try:
      if self.df.empty:
        logger.warning("DataFrame is empty skipping Excel generation.")
        return None

      wb = Workbook()
      ws = wb.active
      ws.title = "Vendors"

      for row_data in dataframe_to_rows(self.df, index=False, header=True):
        ws.append(row_data)

      bold_font = Font(bold=True)
      cold_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
      row_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
      thin_side_h = Side(border_style="thin", color="D9D9D9")
      thin_side_v = Side(border_style="thin", color="808080")

      # Style header
      for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
          cell.font = bold_font
          cell.fill = cold_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h,
                               left=thin_side_v, right=thin_side_v)

      ws.freeze_panes = ws["A2"]

      # Style data rows
      for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
          cell.fill = row_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h,
                               left=thin_side_v, right=thin_side_v)

      col_count = len(self.df.columns)
      if col_count > 0:
        ws.merge_cells(f"A1:{chr(64 + col_count)}1")
        ws["A1"] = f"Vendor Export - {self.company_name}"
        ws["A1"].font = bold_font
        ws["A1"].alignment = Alignment(horizontal="center")

      wb.save(filename)
      return filename

    except Exception as e:
      logger.error(f"export_to_excel() error: {str(e)}")
      return None
