import logging
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

logger = logging.getLogger(__name__)

class ProductServiceExport:
  def __init__(self, products_data, company_name='ERP Inc'):
    self.products_data = products_data
    self.company_name = company_name
    self.formatted_data = []
    self.df = pd.DataFrame()

  def format_data(self):
    try:
      temp = []
      for item in self.products_data:
        try:
          temp.append({
            "ID": item.get("id",""),
            "Name": item.get("item",""),
            "SKU": item.get("sku",""),
            "sale_price": item.get("sale_price",""),
            "purchase_price": item.get("purchase_price",""),
            "Tax": item.get("tax",""),
            "Category": item.get("category",""),
            "Unit": item.get("unit",""),
            "Type": item.get("type",""),
            "Description": item.get("description","")
          })
        except KeyError as e:
          logger.error(f"KeyError in product service: {str(e)}")
        except Exception as e:
          logger.error(f"format_data item exception: {str(e)}")
      self.formatted_data = temp
    except Exception as e:
      logger.error(f"format_data() error: {str(e)}")

  def dataframe(self):
    try:
      if not self.formatted_data:
        logger.warning("No product service data to build dataframe.")
        return pd.DataFrame()
      self.df = pd.DataFrame(self.formatted_data)
      self.df.replace({np.nan:""}, inplace=True)
    except Exception as e:
      logger.error(f"dataframe() error: {str(e)}")
    return self.df

  def export_to_excel(self, filename="product_service_export.xlsx"):
    try:
      if self.df.empty:
        logger.warning("DataFrame is empty skipping Excel creation.")
        return None
      wb = Workbook()
      ws = wb.active
      ws.title = "Products"
      for r in dataframe_to_rows(self.df, index=False, header=True):
        ws.append(r)
      bold_font = Font(bold=True)
      cold_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
      row_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
      thin_side_h = Side(border_style="thin", color="D9D9D9")
      thin_side_v = Side(border_style="thin", color="808080")

      for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
          cell.font = bold_font
          cell.fill = cold_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h,
                               left=thin_side_v, right=thin_side_v)

      ws.freeze_panes = ws["A2"]

      for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
          cell.fill = row_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h,
                               left=thin_side_v, right=thin_side_v)

      cols_count = len(self.df.columns)
      if cols_count>0:
        top_cell = f"A1:{chr(64+cols_count)}1"
        ws.merge_cells(top_cell)
        ws["A1"] = f"Product Service Export - {self.company_name}"
        ws["A1"].font = bold_font
        ws["A1"].alignment = Alignment(horizontal="center")

      wb.save(filename)
      return filename
    except PermissionError as pe:
      logger.error(f"Permission error: {str(pe)}")
      return None
    except Exception as e:
      logger.error(f"export_to_excel() error: {str(e)}")
      return None
