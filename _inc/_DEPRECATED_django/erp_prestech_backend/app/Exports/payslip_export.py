import logging
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

logger = logging.getLogger(__name__)

class PayslipExport:
  def __init__(self, request_data, company_name='ERP Inc'):
    self.request_data = request_data
    self.company_name = company_name
    self.formatted_data = []
    self.df = pd.DataFrame()

  def format_data(self):
    try:
      filter_month = self.request_data.get('filter_month','')
      filter_year = self.request_data.get('filter_year','')
      if not filter_month:
        filter_month = pd.Timestamp.now().shift(-1, 'M').strftime('%m')
      if not filter_year:
        filter_year = pd.Timestamp.now().strftime('%Y')
      formate_month_year = f"{filter_year}-{filter_month}"
      # Here you would normally query your model, e.g. Payslip.objects.filter()
      # We'll build dummy data to simulate:
      dummy_data = [
        {
          "employee_id": "EMP-1001",
          "employee_name": "Alice Worker",
          "basic_salary": "2000",
          "net_salary": "1800",
          "status": "Paid",
          "account_holder_name": "Alice W",
          "account_number": "111122223333",
          "bank_name": "Global Bank",
          "bank_identifier_code": "GB1234",
          "branch_location": "Main City",
          "tax_payer_id": "TX123",
          "month_year": formate_month_year
        },
        {
          "employee_id": "EMP-1002",
          "employee_name": "Bob Worker",
          "basic_salary": "2500",
          "net_salary": "2200",
          "status": "UnPaid",
          "account_holder_name": "Bob W",
          "account_number": "444455556666",
          "bank_name": "Global Bank",
          "bank_identifier_code": "GB5678",
          "branch_location": "Town Branch",
          "tax_payer_id": "TX456",
          "month_year": formate_month_year
        }
      ]
      formatted = []
      for record in dummy_data:
        try:
          formatted.append({
            "EMP ID": record.get("employee_id",""),
            "Name": record.get("employee_name",""),
            "Salary": record.get("basic_salary",""),
            "Net Salary": record.get("net_salary",""),
            "Status": record.get("status",""),
            "Account Holder Name": record.get("account_holder_name",""),
            "Account Number": record.get("account_number",""),
            "Bank Name": record.get("bank_name",""),
            "Bank Identifier Code": record.get("bank_identifier_code",""),
            "Branch Location": record.get("branch_location",""),
            "Tax Payer Id": record.get("tax_payer_id","")
          })
        except KeyError as ke:
          logger.error(f"KeyError formatting payslip: {str(ke)}")
        except Exception as ex:
          logger.error(f"Error in record formatting: {str(ex)}")
      self.formatted_data = formatted
    except Exception as e:
      logger.error(f"format_data() error: {str(e)}")

  def dataframe(self):
    try:
      if not self.formatted_data:
        return pd.DataFrame()
      self.df = pd.DataFrame(self.formatted_data)
      self.df.replace({np.nan:""}, inplace=True)
    except Exception as e:
      logger.error(f"dataframe() error: {str(e)}")
    return self.df

  def export_to_excel(self, filename="payslip_export.xlsx"):
    try:
      if self.df.empty:
        logger.warning("DataFrame is empty skipping Excel generation.")
        return None
      wb = Workbook()
      ws = wb.active
      ws.title = "Payslips"
      for r in dataframe_to_rows(self.df, index=False, header=True):
        ws.append(r)
      bold_font = Font(bold=True)
      cold_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
      row_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
      thin_side_h = Side(border_style="thin", color="D9D9D9")
      thin_side_v = Side(border_style="thin", color="808080")
      for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
          cell.font = bold_font
          cell.fill = cold_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h,
            left=thin_side_v, right=thin_side_v)
      ws.freeze_panes = ws["A2"]
      for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
          cell.fill = row_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h,
            left=thin_side_v, right=thin_side_v)
      cols_count = len(self.df.columns)
      if cols_count>0:
        top_cell = f"A1:{chr(64+cols_count)}1"
        ws.merge_cells(top_cell)
        ws["A1"] = f"Payslip Export - {self.company_name}"
        ws["A1"].font = bold_font
        ws["A1"].alignment = Alignment(horizontal="center")
      wb.save(filename)
      return filename
    except Exception as e:
      logger.error(f"export_to_excel() error: {str(e)}")
      return None
