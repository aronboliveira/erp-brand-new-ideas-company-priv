import logging
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

logger = logging.getLogger(__name__)

class TaskReportExport:
  def __init__(self, project_id, company_name='ERP Inc'):
    self.project_id = project_id
    self.company_name = company_name
    self.tasks_data = []
    self.formatted_data = []
    self.df = pd.DataFrame()

  def format_data(self):
    try:
      temp = []
      # Normally you would query your model with e.g. ProjectTask.objects.filter(...)
      # We'll simulate the data:
      dummy_tasks = [
        {
          "id": 101,
          "title": "Design Phase",
          "description": "Initial project design",
          "start_date": "2024-01-01",
          "end_date": "2024-01-15",
          "priority": "High",
          "assign_to": 10,
          "milestone_id": 5,
          "stage_id": 1
        },
        {
          "id": 102,
          "title": "Implementation",
          "description": "Coding the project",
          "start_date": "2024-01-16",
          "end_date": "2024-02-01",
          "priority": "Medium",
          "assign_to": 11,
          "milestone_id": 6,
          "stage_id": 2
        }
      ]
      # In the real code, you'd filter for tasks using self.project_id and "created_by" from context

      for task in dummy_tasks:
        try:
          # Simulate project_report custom methods
          user_name = self.assign_user(task.get("assign_to", None))
          milestone_name = self.milestone(task.get("milestone_id", None))
          status_name = self.status(task.get("stage_id", None))

          temp.append({
            "ID": task.get("id", ""),
            "Title": task.get("title", ""),
            "Description": task.get("description", ""),
            "Start Date": task.get("start_date", ""),
            "End Date": task.get("end_date", ""),
            "Priority": task.get("priority", ""),
            "Assign To": user_name,
            "Milestone": milestone_name,
            "Status": status_name
          })
        except Exception as exc:
          logger.error(f"Error formatting task: {str(exc)}")

      self.formatted_data = temp
    except Exception as e:
      logger.error(f"format_data() error: {str(e)}")

  def dataframe(self):
    try:
      if not self.formatted_data:
        logger.warning("No formatted data to create DataFrame for TaskReportExport.")
        return pd.DataFrame()
      self.df = pd.DataFrame(self.formatted_data)
      self.df.replace({np.nan: ""}, inplace=True)
    except Exception as e:
      logger.error(f"dataframe() error: {str(e)}")
    return self.df

  def export_to_excel(self, file_name="task_report_export.xlsx"):
    try:
      if self.df.empty:
        logger.warning("DataFrame is empty skipping Excel generation.")
        return None
      wb = Workbook()
      ws = wb.active
      ws.title = "Task Reports"

      for row_data in dataframe_to_rows(self.df, index=False, header=True):
        ws.append(row_data)

      bold_font = Font(bold=True)
      cold_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
      row_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
      thin_side_h = Side(border_style="thin", color="D9D9D9")
      thin_side_v = Side(border_style="thin", color="808080")

      # Header styling
      for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
          cell.font = bold_font
          cell.fill = cold_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h,
                               left=thin_side_v, right=thin_side_v)

      ws.freeze_panes = ws["A2"]

      # Data rows styling
      for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
          cell.fill = row_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h,
                               left=thin_side_v, right=thin_side_v)

      # Merge top cell for a title
      columns_count = len(self.df.columns)
      if columns_count > 0:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns_count)
        ws["A1"] = f"Task Report Export - {self.company_name}"
        ws["A1"].font = bold_font
        ws["A1"].alignment = Alignment(horizontal="center")

      wb.save(file_name)
      return file_name
    except Exception as e:
      logger.error(f"export_to_excel() error: {str(e)}")
      return None

  def headings(self):
    return [
      "ID",
      "Title",
      "Description",
      "Start Date",
      "End Date",
      "Priority",
      "Assign To",
      "Milestone",
      "Status"
    ]

  # Simulated references to project_report static methods
  def assign_user(self, user_id):
    # In real code, might fetch user name from DB
    # We'll stub this out
    if user_id == 10:
      return "Alice"
    elif user_id == 11:
      return "Bob"
    return "Unassigned"

  def milestone(self, milestone_id):
    # Stub for custom method
    if milestone_id == 5:
      return "Design"
    elif milestone_id == 6:
      return "Implementation"
    return "No milestone"

  def status(self, stage_id):
    # Another stub method
    if stage_id == 1:
      return "In Progress"
    elif stage_id == 2:
      return "Completed"
    return "Unknown"
