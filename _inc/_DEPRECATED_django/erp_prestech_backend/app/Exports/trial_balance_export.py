import logging
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

logger = logging.getLogger(__name__)

class TrialBalanceExport:
  def __init__(self, data_dict, start_date, end_date, company_name):
    self.raw_data = data_dict
    self.start_date = start_date
    self.end_date = end_date
    self.company_name = company_name
    self.df = pd.DataFrame()
    self.formatted_data = []

  def format_data(self):
    try:
      temp = []
      total_debit = 0
      total_credit = 0
      for key, accounts_list in self.raw_data.items():
        temp.append({
          "Account Name": "",
          "Account No": "",
          "Debit": "",
          "Credit": ""
        })
        temp.append({
          "Account Name": key,
          "Account No": "",
          "Debit": "",
          "Credit": ""
        })
        for account in accounts_list:
          name = account.get("name", "")
          code = account.get("code", "")
          debit_val = account.get("totalDebit", 0)
          credit_val = account.get("totalCredit", 0)
          temp.append({
            "Account Name": name,
            "Account No": code,
            "Debit": debit_val,
            "Credit": credit_val
          })
          total_debit += debit_val
          total_credit += credit_val
      if temp:
        temp.append({
          "Account Name": "Total",
          "Account No": "",
          "Debit": total_debit,
          "Credit": total_credit
        })
      self.formatted_data = temp
    except Exception as e:
      logger.error(f"format_data() error: {str(e)}")

  def dataframe(self):
    try:
      if not self.formatted_data:
        logger.warning("No formatted data for TrialBalanceExport.")
        return pd.DataFrame()
      self.df = pd.DataFrame(self.formatted_data)
      self.df.replace({np.nan: ""}, inplace=True)
    except Exception as e:
      logger.error(f"dataframe() error: {str(e)}")
    return self.df

  def export_to_excel(self, filename="trial_balance_export.xlsx"):
    try:
      if self.df.empty:
        logger.warning("DataFrame is empty skipping Excel generation.")
        return None
      wb = Workbook()
      ws = wb.active
      ws.title = "Trial Balance"

      for row_vals in dataframe_to_rows(self.df, index=False, header=True):
        ws.append(row_vals)

      bold_font = Font(bold=True)
      cold_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
      row_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
      thin_side_h = Side(border_style="thin", color="D9D9D9")
      thin_side_v = Side(border_style="thin", color="808080")

      # Style header
      for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
          cell.font = bold_font
          cell.fill = cold_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h,
                               left=thin_side_v, right=thin_side_v)

      ws.freeze_panes = ws["A2"]

      # Data row styling
      for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
          cell.fill = row_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h,
                               left=thin_side_v, right=thin_side_v)

      # Merge top cells for a title
      max_cols = len(self.df.columns)
      if max_cols > 0:
        ws.merge_cells(f"A1:{chr(64 + max_cols)}1")
        ws.merge_cells(f"A2:{chr(64 + max_cols)}2")
        ws.merge_cells(f"A3:{chr(64 + max_cols)}3")
        ws["A1"] = f"Trial Balance - {self.company_name}"
        ws["A2"] = f"Print Out Date: {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')}"
        ws["A3"] = f"Date: {self.start_date} - {self.end_date}"
        ws["A1"].font = bold_font
        ws["A2"].font = bold_font
        ws["A3"].font = bold_font
        for row_label in ["A1", "A2", "A3"]:
          ws[row_label].alignment = Alignment(horizontal="center")

      # Bold last row
      last_row = ws.max_row
      for cell in ws[f"A{last_row}:Z{last_row}"][0]:
        cell.font = Font(bold=True)

      # Additional logic: Bold certain row if it equals Assets, Income, etc.
      # We'll replicate the Maatwebsite behavior:
      for idx, rowdata in enumerate(self.df.values.tolist(), start=2):
        # 2 offset because row 1 is header
        acct_name = rowdata[0]
        if acct_name in ["Assets", "Income", "Costs of Goods Sold", "Expenses", "Liabilities", "Equity"]:
          row_index = idx
          for cell in ws[f"A{row_index}:D{row_index}"][0]:
            cell.font = Font(bold=True)

      wb.save(filename)
      return filename
    except Exception as e:
      logger.error(f"export_to_excel() error: {str(e)}")
      return None
