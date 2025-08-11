import logging
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import inch
import json
# from app.Exports.views.balance_sheet_export_api_view import BalanceSheetExportAPIView
# from django.urls import path
logger = logging.getLogger(__name__)

class AccountStatementExport:
  def __init__(self, data_sets, sheet_names, start_date, end_date, company_name, theme=None):
    self.data_sets = data_sets
    self.sheet_names = sheet_names
    self.start_date = start_date
    self.end_date = end_date
    self.company_name = company_name
    self.theme = theme if theme else self.default_theme()
    self.workbook = Workbook()

  def default_theme(self):
    return {
      'header_color': '003366',
      'header_font': 'Arial',
      'header_size': 12,
      'row_fill': 'D9E1F2',
      'border_color': '808080',
      'logo_path': None
    }

  def dataframe_from_data(self, data):
    try:
      if not data:
        logger.warning("Empty data set received.")
        return pd.DataFrame()
      df = pd.DataFrame(data)
      df.replace({np.nan: ''}, inplace=True)
      return df
    except Exception as e:
      logger.error(f"dataframe_from_data() error: {str(e)}")
      return pd.DataFrame()

  def export_to_excel(self, file_name="multi_sheet.xlsx"):
    try:
      wb = self.workbook
      first_sheet = True
      for idx, data in enumerate(self.data_sets):
        df = self.dataframe_from_data(data)
        if df.empty:
          logger.warning(f"Sheet {self.sheet_names[idx]} is empty skipping.")
          continue
        if first_sheet:
          ws = wb.active
          ws.title = self.sheet_names[idx]
          first_sheet = False
        else:
          ws = wb.create_sheet(title=self.sheet_names[idx])
        for r in dataframe_to_rows(df, index=False, header=True):
          ws.append(r)
        bold_font = Font(bold=True, name=self.theme['header_font'], size=self.theme['header_size'])
        cold_fill = PatternFill(start_color=self.theme['header_color'], end_color=self.theme['header_color'], fill_type='solid')
        row_fill = PatternFill(start_color=self.theme['row_fill'], end_color=self.theme['row_fill'], fill_type='solid')
        thin_border_horizontal = Side(border_style="thin", color="D9D9D9")
        thin_border_vertical = Side(border_style="thin", color=self.theme['border_color'])
        for col in ws.iter_cols(min_row=1, max_row=1):
          for cell in col:
            cell.font = bold_font
            cell.fill = cold_fill
            cell.border = Border(top=thin_border_horizontal, bottom=thin_border_horizontal,
                                 left=thin_border_vertical, right=thin_border_vertical)
        ws.freeze_panes = ws['A2']
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
          for cell in row:
            cell.fill = row_fill
            cell.border = Border(top=thin_border_horizontal, bottom=thin_border_horizontal,
                                 left=thin_border_vertical, right=thin_border_vertical)
        ws.merge_cells('A1:E1')
        ws.merge_cells('A2:E2')
        ws.merge_cells('A3:E3')
        ws['A1'] = f"{self.sheet_names[idx]} - {self.company_name}"
        ws['A2'] = f"Print Out Date: {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')}"
        ws['A3'] = f"Date: {self.start_date} - {self.end_date}"
        ws['A1'].font = bold_font
        ws['A2'].font = bold_font
        ws['A3'].font = bold_font
        ws['A1'].alignment = Alignment(horizontal="center")
        ws['A2'].alignment = Alignment(horizontal="center")
        ws['A3'].alignment = Alignment(horizontal="center")
        for col_idx, column_cells in enumerate(ws.columns, start=1):
          max_length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)
          adjusted_width = (max_length + 2)
          ws.column_dimensions[chr(64 + col_idx)].width = adjusted_width if adjusted_width < 50 else 50
    except Exception as e:
      logger.error(f"export_to_excel() error: {str(e)}")
    try:
      wb.save(file_name)
    except Exception as e:
      logger.error(f"Workbook save error: {str(e)}")

  def export_to_pdf(self, file_name="multi_sheet.pdf"):
    try:
      c = canvas.Canvas(file_name, pagesize=A4)
      width, height = A4
      if self.theme['logo_path']:
        try:
          c.drawImage(self.theme['logo_path'], inch, height - inch * 1.5, width=1 * inch, height=1 * inch)
        except Exception as e:
          logger.error(f"Error loading logo: {str(e)}")
      y = height - 2 * inch
      for idx, data in enumerate(self.data_sets):
        c.setFont(self.theme['header_font'], self.theme['header_size'])
        c.drawString(inch, y, f"{self.sheet_names[idx]} - {self.company_name}")
        y -= 20
        df = self.dataframe_from_data(data)
        if df.empty:
          c.drawString(inch, y, "No data available.")
          y -= 20
          continue
        for col in df.columns:
          c.drawString(inch, y, str(col))
          y -= 15
        y -= 10
        for _, row in df.iterrows():
          row_text = " | ".join([str(val) for val in row.tolist()])
          c.drawString(inch, y, row_text)
          y -= 15
          if y < inch:
            c.showPage()
            y = height - inch
      c.save()
    except Exception as e:
      logger.error(f"export_to_pdf() error: {str(e)}")

  def ask_to_save_json(self):
    try:
      save_input = input("Do you want to save the results as JSON? (y/n): ")
      if save_input.strip().lower() == 'y':
        result_dict = {sheet_name: data for sheet_name, data in zip(self.sheet_names, self.data_sets)}
        with open('multi_sheet.json', 'w', encoding='utf-8') as f:
          json.dump(result_dict, f, ensure_ascii=False, indent=2)
    except Exception as e:
      logger.error(f"ask_to_save_json() error: {str(e)}")

  def run_exports(self, file_type='excel'):
    try:
      if file_type == 'pdf':
        self.export_to_pdf()
      else:
        self.export_to_excel()
      self.ask_to_save_json()
    except Exception as e:
      logger.error(f"run_exports() error: {str(e)}")
      
