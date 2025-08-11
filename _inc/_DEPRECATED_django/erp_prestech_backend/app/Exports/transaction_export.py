import logging
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

logger = logging.getLogger(__name__)

class TransactionExport:
  def __init__(self, transactions_data, company_name='ERP Inc'):
    self.transactions_data = transactions_data
    self.company_name = company_name
    self.formatted_data = []
    self.df = pd.DataFrame()

  def format_data(self):
    try:
      temp = []
      # In an actual app, you'd fetch from Transaction model. We'll just parse the input array.
      for tx in self.transactions_data:
        try:
          account_info = self.accounts(tx.get("account",""))
          temp.append({
            "Transaction Id": tx.get("transaction_id",""),
            "Account": account_info,
            "Type": tx.get("type",""),
            "Amount": tx.get("amount",0),
            "Description": tx.get("description",""),
            "Date": tx.get("date",""),
            "Category": tx.get("category","")
          })
        except Exception as e:
          logger.error(f"Error formatting transaction: {str(e)}")
      self.formatted_data = temp
    except Exception as e:
      logger.error(f"format_data() error: {str(e)}")

  def accounts(self, account_field):
    # Simulate Transaction::accounts($Transaction->account)
    # In real usage, this might do a DB lookup or reference a static method
    # We'll just do a placeholder transform
    return f"Account: {account_field}"

  def dataframe(self):
    try:
      if not self.formatted_data:
        logger.warning("No data to build DataFrame for TransactionExport.")
        return pd.DataFrame()
      self.df = pd.DataFrame(self.formatted_data)
      self.df.replace({np.nan:""}, inplace=True)
    except Exception as e:
      logger.error(f"dataframe() error: {str(e)}")
    return self.df

  def export_to_excel(self, filename="transaction_export.xlsx"):
    try:
      if self.df.empty:
        logger.warning("DataFrame is empty skipping Excel generation.")
        return None
      wb = Workbook()
      ws = wb.active
      ws.title = "Transactions"
      for row_data in dataframe_to_rows(self.df, index=False, header=True):
        ws.append(row_data)
      bold_font = Font(bold=True)
      cold_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
      row_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
      thin_side_h = Side(border_style="thin", color="D9D9D9")
      thin_side_v = Side(border_style="thin", color="808080")

      # Style headers
      for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
          cell.font = bold_font
          cell.fill = cold_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h,
                               left=thin_side_v, right=thin_side_v)

      ws.freeze_panes = ws["A2"]

      # Style data rows
      for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
          cell.fill = row_fill
          cell.border = Border(top=thin_side_h, bottom=thin_side_h,
                               left=thin_side_v, right=thin_side_v)

      cols_count = len(self.df.columns)
      if cols_count > 0:
        top_cell = f"A1:{chr(64 + cols_count)}1"
        ws.merge_cells(top_cell)
        ws["A1"] = f"Transaction Export - {self.company_name}"
        ws["A1"].font = bold_font
        ws["A1"].alignment = Alignment(horizontal="center")

      wb.save(filename)
      return filename
    except Exception as e:
      logger.error(f"export_to_excel() error: {str(e)}")
      return None
